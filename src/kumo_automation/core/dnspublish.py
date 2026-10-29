"""cPanel DNS publishing — add / delete / update / verify zone records.

Adapts Suresh's cPanel bulk tools (add_dns.py + flush_dns.py) into the app and
extends them beyond A records to TXT (SPF/DKIM/DMARC) and MX, so the whole
generated record set publishes from one place. cPanel API2 ZoneEdit, API-token
auth (Authorization: cpanel <user>:<token>).

Design:
  * planning functions (plan_add / plan_delete_by_range / record_params) are PURE
    — no network — so they are unit-tested without a cPanel.
  * CpanelClient does the I/O; inject a fake .call in tests.
Safety kept from Suresh: dry-run separate from apply, conflict surfaced not
overwritten, backup before delete, per-line edit for update. PTR is out of scope
(reverse DNS is set at the IP provider, not cPanel).
"""
from __future__ import annotations

import ipaddress
import json
import ssl
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Optional

from .models import DnsRecord

LogFn = Callable[[str], None]


# --------------------------------------------------------------------------- #
#  Credentials
# --------------------------------------------------------------------------- #
@dataclass
class CpanelCreds:
    host: str
    user: str
    token: str

    def normalized_host(self) -> str:
        host = self.host
        if not host.startswith(("http://", "https://")):
            host = "https://" + host
        if ":" not in host.split("//", 1)[1]:
            host += ":2083"
        return host.rstrip("/")


# --------------------------------------------------------------------------- #
#  Record -> cPanel params  (pure)
# --------------------------------------------------------------------------- #
def record_params(rec: DnsRecord) -> dict:
    """Map a DnsRecord to add/edit_zone_record params, by type."""
    name = rec.name.rstrip(".") + "."
    p: dict = {"name": name, "type": rec.type.upper(), "ttl": rec.ttl, "class": "IN"}
    t = rec.type.upper()
    if t in ("A", "AAAA"):
        p["address"] = rec.value
    elif t == "TXT":
        p["txtdata"] = rec.value
    elif t == "CNAME":
        p["cname"] = rec.value.rstrip(".") + "."
    elif t == "MX":
        pri, _, ex = rec.value.partition(" ")
        p["priority"] = pri.strip() or "10"
        p["exchange"] = ex.strip().rstrip(".") + "."
    else:
        p["address"] = rec.value
    return p


def _norm(value: str) -> str:
    """Normalize a TXT/record value for comparison (strip quotes + whitespace)."""
    return "".join(value.split()).strip('"').lower()


# --------------------------------------------------------------------------- #
#  Planning  (pure — operate on a fetched zone snapshot)
# --------------------------------------------------------------------------- #
def plan_add(existing: list[dict], desired: list[DnsRecord]):
    """Return (todo, already, conflicts).

    existing: normalized zone records [{name,type,value,line,ttl}].
    already : desired record already present with the same value.
    conflict: same name+type present with a DIFFERENT value (adding would create
              a second record) — surfaced, still added, never silently replaced.
    """
    idx: dict[tuple[str, str], list[dict]] = {}
    for e in existing:
        idx.setdefault((e["name"].rstrip(".").lower(), e["type"].upper()), []).append(e)
    todo, already, conflicts = [], [], []
    for r in desired:
        haves = idx.get((r.name.rstrip(".").lower(), r.type.upper()), [])
        vals = {_norm(h["value"]) for h in haves}
        if _norm(r.value) in vals:
            already.append(r)
        elif haves:
            conflicts.append((r, sorted(h["value"] for h in haves)))
            todo.append(r)
        else:
            todo.append(r)
    return todo, already, conflicts


def plan_delete_by_range(existing: list[dict], cidrs: list[str]) -> list[dict]:
    """Existing records whose value is an IP inside any of the CIDRs (flush)."""
    nets = [ipaddress.ip_network(c.strip(), strict=False) for c in cidrs if c.strip()]
    matched = []
    for e in existing:
        try:
            ip = ipaddress.ip_address(e["value"])
        except ValueError:
            continue
        if any(ip in n for n in nets):
            matched.append(e)
    return matched


def find_line(existing: list[dict], name: str, rtype: str,
              value: Optional[str] = None) -> Optional[int]:
    """Zone line number of a record matching name+type (+value), for edit/remove."""
    for e in existing:
        if (e["name"].rstrip(".").lower() == name.rstrip(".").lower()
                and e["type"].upper() == rtype.upper()
                and (value is None or _norm(e["value"]) == _norm(value))):
            return e.get("line")
    return None


# --------------------------------------------------------------------------- #
#  cPanel client  (I/O)
# --------------------------------------------------------------------------- #
class CpanelError(RuntimeError):
    pass


class CpanelClient:
    def __init__(self, creds: CpanelCreds, insecure: bool = False,
                 opener: Optional[Callable[[str, dict], str]] = None):
        self.creds = creds
        self.host = creds.normalized_host()
        self.ctx = ssl._create_unverified_context() if insecure else None
        self._opener = opener               # test seam: (url, headers) -> body text

    def call(self, module: str, func: str, **params) -> dict:
        query = {
            "cpanel_jsonapi_user": self.creds.user,
            "cpanel_jsonapi_apiversion": "2",
            "cpanel_jsonapi_module": module,
            "cpanel_jsonapi_func": func,
        }
        query.update(params)
        url = f"{self.host}/json-api/cpanel?{urllib.parse.urlencode(query)}"
        headers = {"Authorization": f"cpanel {self.creds.user}:{self.creds.token}"}
        if self._opener:
            return json.loads(self._opener(url, headers))
        req = urllib.request.Request(url, headers=headers)
        try:
            with urllib.request.urlopen(req, timeout=60, context=self.ctx) as resp:
                return json.loads(resp.read().decode("utf-8", "replace"))
        except urllib.error.HTTPError as exc:
            raise CpanelError(f"HTTP {exc.code}: {exc.read().decode('utf-8','replace')[:300]}") from None
        except urllib.error.URLError as exc:
            raise CpanelError(f"Cannot reach {self.host}: {exc.reason}") from None

    def fetch_zone(self, domain: str, rtype: Optional[str] = None) -> list[dict]:
        params = {"domain": domain}
        if rtype:
            params["type"] = rtype
        block = self.call("ZoneEdit", "fetchzone_records", **params).get("cpanelresult", {})
        if block.get("error"):
            raise CpanelError(str(block["error"]))
        out = []
        for e in block.get("data") or []:
            value = (e.get("address") or e.get("record") or e.get("txtdata")
                     or e.get("cname") or e.get("exchange") or "")
            out.append({
                "name": (e.get("name") or "").rstrip(".") + ".",
                "type": (e.get("type") or "A").upper(),
                "value": str(value).strip(),
                "ttl": e.get("ttl"),
                "line": e.get("Line") or e.get("line"),
            })
        return out

    def _status_ok(self, result: dict) -> tuple[bool, str]:
        block = result.get("cpanelresult", {})
        if block.get("error"):
            return False, str(block["error"])
        rows = block.get("data") or []
        inner = rows[0].get("result", rows[0]) if rows else {}
        if str(inner.get("status")) == "1":
            return True, inner.get("statusmsg", "")
        return False, inner.get("statusmsg") or json.dumps(block)[:200]

    def add(self, domain: str, rec: DnsRecord) -> tuple[bool, str]:
        return self._status_ok(self.call("ZoneEdit", "add_zone_record",
                                          domain=domain, **record_params(rec)))

    def remove(self, domain: str, line: int) -> tuple[bool, str]:
        return self._status_ok(self.call("ZoneEdit", "remove_zone_record",
                                          domain=domain, line=line))

    def edit(self, domain: str, line: int, rec: DnsRecord) -> tuple[bool, str]:
        return self._status_ok(self.call("ZoneEdit", "edit_zone_record",
                                          domain=domain, line=line, **record_params(rec)))


# --------------------------------------------------------------------------- #
#  High-level operations (add / delete / update / verify)
# --------------------------------------------------------------------------- #
# PTR is not a cPanel record — reverse DNS lives at the IP provider.
def _publishable(records: list[DnsRecord]) -> list[DnsRecord]:
    return [r for r in records if r.type.upper() != "PTR"]


def publish(client: CpanelClient, domain: str, records: list[DnsRecord],
            *, dry_run: bool = True, log: LogFn = print) -> dict:
    """ADD every publishable record to `domain`. Dry-run reports the plan only."""
    desired = _publishable(records)
    existing = client.fetch_zone(domain)
    todo, already, conflicts = plan_add(existing, desired)
    log(f"{domain}: {len(existing)} existing | to add {len(todo)}, "
        f"skip {len(already)} dup, {len(conflicts)} conflict")
    for r, have in conflicts:
        log(f"  ! {r.name} {r.type} already = {have} — would add {r.value} as additional")
    if dry_run:
        return {"added": 0, "todo": len(todo), "already": len(already),
                "conflicts": len(conflicts), "failed": [], "dry_run": True}
    added, failed = 0, []
    for i, r in enumerate(todo, 1):
        ok, msg = client.add(domain, r)
        if ok:
            added += 1
            log(f"[{i}/{len(todo)}] OK   {r.type:5} {r.name} -> {r.value}")
        else:
            failed.append((r.name, r.type, msg))
            log(f"[{i}/{len(todo)}] FAIL {r.type:5} {r.name}: {msg}")
    return {"added": added, "todo": len(todo), "already": len(already),
            "conflicts": len(conflicts), "failed": failed, "dry_run": False}


def delete_range(client: CpanelClient, domain: str, cidrs: list[str],
                 *, confirm: bool = False, backup_path: Optional[str] = None,
                 log: LogFn = print) -> dict:
    """DELETE records whose IP falls in `cidrs`. Backs up first; report-only
    unless confirm=True (Suresh's flush behaviour)."""
    existing = client.fetch_zone(domain)
    matched = plan_delete_by_range(existing, cidrs)
    log(f"{domain}: {len(matched)} record(s) match {', '.join(cidrs)}")
    if backup_path and matched:
        Path(backup_path).write_text(json.dumps(matched, indent=2), encoding="utf-8")
        log(f"  backup -> {backup_path}")
    if not confirm:
        return {"matched": len(matched), "deleted": 0, "dry_run": True}
    if any(not m.get("line") for m in matched):
        raise CpanelError("some matched records have no line number; refusing to delete")
    deleted, failed = 0, []
    # descending line order — removing a line renumbers everything below it
    for m in sorted(matched, key=lambda x: int(x["line"]), reverse=True):
        ok, msg = client.remove(domain, int(m["line"]))
        if ok:
            deleted += 1
            log(f"  deleted {m['name']} -> {m['value']}")
        else:
            failed.append((m["name"], msg))
            log(f"  FAILED  {m['name']}: {msg}")
    return {"matched": len(matched), "deleted": deleted, "failed": failed, "dry_run": False}


def update_record(client: CpanelClient, domain: str, rec: DnsRecord,
                  *, dry_run: bool = True, log: LogFn = print) -> dict:
    """UPDATE a record in place (edit_zone_record on its existing line)."""
    existing = client.fetch_zone(domain)
    line = find_line(existing, rec.name, rec.type)
    if line is None:
        log(f"{domain}: no existing {rec.type} for {rec.name} to edit — use add instead")
        return {"ok": False, "reason": "not found; use add"}
    if dry_run:
        log(f"{domain}: would edit line {line}: {rec.name} {rec.type} -> {rec.value}")
        return {"ok": True, "line": line, "dry_run": True}
    ok, msg = client.edit(domain, int(line), rec)
    log(f"{domain}: edit line {line} {rec.name} -> {rec.value}: {'OK' if ok else msg}")
    return {"ok": ok, "line": line, "msg": msg}


def verify(client: CpanelClient, domain: str, expected: list[DnsRecord]) -> dict:
    """Diff expected records against the live zone."""
    existing = client.fetch_zone(domain)
    todo, already, conflicts = plan_add(existing, _publishable(expected))
    return {"published": len(already), "missing": len(todo),
            "conflicts": len(conflicts), "missing_records": todo}


# --------------------------------------------------------------------------- #
#  Excel intake + domain grouping + batch publish
# --------------------------------------------------------------------------- #
def domain_of(name: str, known: Optional[list[str]] = None) -> str:
    """Zone a record belongs to. Prefers matching a KNOWN domain (handles any
    depth incl. multi-label TLDs); else falls back to the last two labels.
    Strips DKIM/DMARC prefixes naturally (last-two-labels of _dmarc.d.com = d.com)."""
    n = name.rstrip(".").lower()
    if known:
        for d in sorted((k.lower() for k in known), key=len, reverse=True):
            if n == d or n.endswith("." + d):
                return d
    labels = n.split(".")
    return ".".join(labels[-2:]) if len(labels) >= 2 else n


def group_by_domain(records: list[DnsRecord],
                    known: Optional[list[str]] = None) -> dict[str, list[DnsRecord]]:
    grouped: dict[str, list[DnsRecord]] = {}
    for r in records:
        grouped.setdefault(domain_of(r.name, known), []).append(r)
    return grouped


def parse_records_xlsx(path: str) -> list[DnsRecord]:
    """Read a DNS sheet into DnsRecord[]. Matches the app's Download-Center export
    (Type, Host, Value, TTL, …) and also Suresh's layout (Name, TTL, Type, Record);
    columns are found by header name in any order. PTR rows are kept (grouping /
    publish drop them where appropriate)."""
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(c or "").strip().lower() for c in rows[0]]

    def col(*names):
        for nm in names:
            if nm in header:
                return header.index(nm)
        return None

    i_type = col("type")
    i_name = col("host", "name", "hostname")
    i_value = col("value", "record", "address", "data")
    i_ttl = col("ttl")
    if i_name is None or i_value is None:
        raise CpanelError(f"could not find Host/Value columns in header: {header}")

    out: list[DnsRecord] = []
    for row in rows[1:]:
        if not row or all(c is None for c in row):
            continue
        name = str(row[i_name] or "").strip()
        value = str(row[i_value] or "").strip()
        if not name or not value:
            continue
        rtype = (str(row[i_type]).strip().upper() if i_type is not None and row[i_type] else "A")
        ttl = int(row[i_ttl]) if i_ttl is not None and str(row[i_ttl] or "").strip().isdigit() else 3600
        out.append(DnsRecord(type=rtype, name=name, value=value, ttl=ttl))
    return out


def routes_from_accounts(accounts, insecure: bool = False) -> dict[str, "CpanelClient"]:
    """Build {domain: CpanelClient} from cPanel accounts, using keyring tokens.
    `accounts` = iterable with .host/.user/.domains. Accounts with no stored
    token are skipped (their domains stay unrouted)."""
    routes: dict[str, CpanelClient] = {}
    for acc in accounts:
        creds = creds_from_keyring(acc.host, acc.user)
        if not creds:
            continue
        client = CpanelClient(creds, insecure=insecure)
        for d in acc.domains:
            routes[d.lower()] = client
    return routes


def publish_many(routes: dict[str, "CpanelClient"], records: list[DnsRecord],
                 *, dry_run: bool = True, log: LogFn = print) -> dict:
    """Publish a mixed-domain record set: group by domain, route each to its
    account's client, publish. Returns {results:{domain:...}, unrouted:[domain]}."""
    grouped = group_by_domain(_publishable(records), known=list(routes.keys()))
    results: dict[str, dict] = {}
    unrouted: list[str] = []
    for domain, recs in grouped.items():
        client = routes.get(domain.lower())
        if client is None:
            unrouted.append(domain)
            log(f"{domain}: NO cPanel account configured — skipped ({len(recs)} record(s))")
            continue
        results[domain] = publish(client, domain, recs, dry_run=dry_run, log=log)
    return {"results": results, "unrouted": unrouted}


# --------------------------------------------------------------------------- #
#  Credentials — keyring first (per cPanel account), file fallback (Suresh's)
# --------------------------------------------------------------------------- #
def _key(host: str, user: str) -> str:
    return f"cpanel::{host}::{user}"


def store_token(host: str, user: str, token: str) -> None:
    from ..services import secrets
    secrets.set_secret(_key(host, user), token)


def creds_from_keyring(host: str, user: str) -> Optional[CpanelCreds]:
    from ..services import secrets
    token = secrets.get_secret(_key(host, user))
    return CpanelCreds(host=host, user=user, token=token) if token else None


def creds_from_file(path: str) -> CpanelCreds:
    data: dict[str, str] = {}
    for line in Path(path).read_text(encoding="utf-8-sig").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, _, v = line.partition("=")
        data[k.strip().upper()] = v.strip().strip('"').strip("'")
    missing = [k for k in ("HOST", "USER", "TOKEN") if not data.get(k)]
    if missing:
        raise CpanelError(f"creds file missing: {', '.join(missing)}")
    return CpanelCreds(host=data["HOST"], user=data["USER"], token=data["TOKEN"])
